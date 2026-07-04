"""Dump the IK India Program Master Index docx (paragraphs + all tables) to a
clean text file used as the gap-analysis baseline."""
import docx, os

DOC = r"C:\Users\Admin\Downloads\IK India - Program Master Index (Live).docx"
OUT = os.path.join(os.path.dirname(__file__), "curriculum_raw.txt")

d = docx.Document(DOC)
lines = []
lines.append("===== PARAGRAPHS =====")
for p in d.paragraphs:
    t = p.text.strip()
    if t:
        lines.append(t)
for ti, tbl in enumerate(d.tables):
    lines.append(f"\n===== TABLE {ti+1} ({len(tbl.rows)}x{len(tbl.columns)}) =====")
    for row in tbl.rows:
        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
        if any(cells):
            lines.append(" | ".join(cells))

txt = "\n".join(lines)
# normalise the mojibake dashes/bullets seen in the dump
txt = txt.replace("�", "-")
open(OUT, "w", encoding="utf-8").write(txt)
print("Wrote", OUT, "chars:", len(txt), "| paragraphs:", len(d.paragraphs), "| tables:", len(d.tables))
