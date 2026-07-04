"""Build CSV + multi-sheet XLSX for the emerging-roles gap analysis.
The narrative gap-analysis .md is authored separately (Claude synthesis)."""
import os, json, csv, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", line_buffering=True)
OUT = os.path.join(os.path.dirname(__file__), "output")

rows = json.load(open(os.path.join(OUT, "emerging_rows.json"), encoding="utf-8"))
agg = json.load(open(os.path.join(OUT, "emerging_aggregates.json"), encoding="utf-8"))
man = json.load(open(os.path.join(OUT, "emerging_manifest.json"), encoding="utf-8"))
import emerging_config as E
for r in rows:
    r.setdefault("skills", E.extract_skills(r.get("desc", "")))

# ---- CSV ----
cpath = os.path.join(OUT, "emerging_ai_roles_india.csv")
with open(cpath, "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(["Source", "Job Title", "Canonical", "Emerging Bucket", "Seniority Bucket",
                "Lenses", "Company", "Company Tier", "Salary", "Location", "Posted Date",
                "Window", "Skills/Tools", "Query", "URL"])
    for r in rows:
        w.writerow([r.get("source", ""), r.get("title", ""), r.get("canonical", ""),
                    r.get("emerging_bucket", "") or "", r.get("seniority_bucket", "") or "",
                    ";".join(r.get("lenses", [])), r.get("company", ""), r.get("company_tier", ""),
                    r.get("salary_norm", "") or r.get("salary", ""), r.get("location", ""),
                    r.get("posted_date", ""), r.get("window", ""),
                    ";".join(r.get("skills", [])), r.get("query", ""), r.get("url", "")])
print("Wrote", cpath)

# ---- XLSX ----
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    HF = PatternFill("solid", fgColor="1F4E78"); HFONT = Font(color="FFFFFF", bold=True)

    def sheet(name, headers, data):
        ws = wb.create_sheet(name); ws.append(headers)
        for c in range(1, len(headers) + 1):
            ws.cell(row=1, column=c).fill = HF; ws.cell(row=1, column=c).font = HFONT
            ws.cell(row=1, column=c).alignment = Alignment(horizontal="center")
        for d in data:
            ws.append(d)
        ws.freeze_panes = "A2"
        for col in ws.columns:
            L = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(L + 2, 60)
        return ws

    s = wb.active; s.title = "Summary"
    s["A1"] = "Emerging AI Roles -> IK Curriculum Gap Analysis (India, last 2 months)"
    s["A1"].font = Font(bold=True, size=14)
    g = man["governor"]
    rowsout = [
        ["Window", f'{man["window"]["start"]} to {man["window"]["report_date"]}'],
        ["AI rows analyzed", agg["n_rows"]],
        ["In 2-month window", agg["n_in_window"]],
        ["Sources", ", ".join(f"{k}:{v}" for k, v in agg["by_source"].items())],
        ["Lens counts", f'GenAI-emerging {agg["lens_counts"]["L1_genai_emerging"]} | '
                        f'seniority {agg["lens_counts"]["L4_seniority_hybrid"]} | '
                        f'novel {agg["lens_counts"]["novel_unmapped"]}'],
        ["Spend (USD)", round(g["session_spent"], 4)],
        ["Final usage / cap", f'{round(g["final_usage"],4)} / 5.00'],
    ]
    for i, (a, b) in enumerate(rowsout, start=3):
        s.cell(row=i, column=1, value=a); s.cell(row=i, column=2, value=b)

    sheet("Emerging_Roles (L1)", ["Emerging Role Bucket", "Openings", "In-Window",
                                  "IK Coverage (auto)", "Top Companies", "Sample Titles"],
          [[b["bucket"], b["count"], b["in_window"], b["coverage_guess"],
            ", ".join(b["companies"][:5]), " | ".join(b["samples"])] for b in agg["emerging_buckets"]])
    sheet("Seniority_Hybrid (L4)", ["Role Bucket", "Openings", "Top Companies", "Sample Titles"],
          [[b["bucket"], b["count"], ", ".join(b["companies"][:5]), " | ".join(b["samples"])]
           for b in agg["seniority_buckets"]])
    sheet("Novel_Unmapped (L2)", ["Canonical Title", "Count", "Sample Raw Title", "Companies"],
          [[b["canonical"], b["count"], b["sample"], ", ".join(b["companies"])] for b in agg["novel_titles"]])
    sheet("Top_Titles", ["Canonical Title", "Openings"],
          [[t["title"], t["count"]] for t in agg["top_canonical"]])
    sheet("Skills_Tools_Demand", ["Skill / Tool", "Mentions", "% of rows", "In IK Curriculum"],
          [[s_["skill"], s_["count"], s_["pct"], "YES" if s_["in_curriculum"] else "NO"]
           for s_ in agg["skills_demand"]])
    sheet("All_Rows", ["Source", "Title", "Emerging Bucket", "Seniority", "Company", "Tier",
                       "Salary", "Posted", "Window", "Skills", "URL"],
          [[r.get("source"), r.get("title"), r.get("emerging_bucket") or "", r.get("seniority_bucket") or "",
            r.get("company"), r.get("company_tier"), r.get("salary_norm") or r.get("salary"),
            r.get("posted_date"), r.get("window"), ";".join(r.get("skills", [])), r.get("url")] for r in rows])
    sheet("Cost_Governance", ["Item", "Value"],
          [["Account plan", "FREE $5 cap"], ["Hard ceiling", "$2.50"], ["Soft ceiling", "$2.20"],
           ["Spend this run", f'${round(g["session_spent"],4)}'],
           ["Final usage", f'${round(g["final_usage"],4)} / $5.00'],
           ["Sources", "Naukri + LinkedIn only"]])

    xp = os.path.join(OUT, "emerging_titles_ranked.xlsx")
    wb.save(xp); print("Wrote", xp)
except Exception as e:
    print("XLSX error:", e)
